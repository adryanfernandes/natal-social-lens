import os
from openpyxl import load_workbook

base = r'C:\git\painel_cadunico\natal-social-lens\src\data'
for name in sorted(os.listdir(base)):
    if not name.lower().endswith('.xlsx'):
        continue
    path = os.path.join(base, name)
    print(f'FILE: {name}')
    wb = load_workbook(path, data_only=True)
    print('SHEETS:', wb.sheetnames)
    for ws in wb.worksheets[:2]:
        print(f'--- SHEET: {ws.title} rows={ws.max_row} cols={ws.max_column}')
        for row in ws.iter_rows(min_row=1, max_row=min(5, ws.max_row), values_only=True):
            print(row)
