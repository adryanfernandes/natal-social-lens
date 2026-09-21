import json
import os
import re
import time
import unicodedata
from collections import Counter
from datetime import date, datetime

import openpyxl
import requests
import truststore


SOURCE_FILE = "2026_BDTrabalhado_10Abril_Karine.xlsx"
DICTIONARY_FILE = "dicionariotudo.xlsx"
CUBE_SOURCE = "dashboard_cube_v4"
BATCH_SIZE = 100

OFFICIAL_NEIGHBORHOODS = {
    "ALECRIM": "Alecrim",
    "AREIA PRETA": "Areia Preta",
    "BARRO VERMELHO": "Barro Vermelho",
    "BOM PASTOR": "Bom Pastor",
    "CANDELARIA": "Candelaria",
    "CAPIM MACIO": "Capim Macio",
    "CIDADE ALTA": "Cidade Alta",
    "CIDADE DA ESPERANCA": "Cidade da Esperanca",
    "CIDADE NOVA": "Cidade Nova",
    "DIX SEPT ROSADO": "Dix-Sept Rosado",
    "FELIPE CAMARAO": "Filipe Camarao",
    "GUARAPES": "Guarapes",
    "IGAPO": "Igapo",
    "JARDIM PROGRESSO": "Jardim Progresso",
    "LAGOA AZUL": "Lagoa Azul",
    "LAGOA NOVA": "Lagoa Nova",
    "LAGOA SECA": "Lagoa Seca",
    "MAE LUIZA": "Mae Luiza",
    "MORRO BRANCO": "Morro Branco",
    "NEOPOLIS": "Neopolis",
    "NORDESTE": "Nordeste",
    "NOSSA SENHORA DA APRESENTACAO": "Nossa Senhora da Apresentacao",
    "NOSSA SENHORA DE NAZARE": "Nossa Senhora de Nazare",
    "NOVA DESCOBERTA": "Nova Descoberta",
    "PAJUCARA": "Pajucara",
    "PARQUE DAS COLINAS": "Parque das Colinas",
    "PETROPOLIS": "Petropolis",
    "PITIMBU": "Pitimbu",
    "PLANALTO": "Planalto",
    "PONTA NEGRA": "Ponta Negra",
    "POTENGI": "Potengi",
    "PRAIA DO MEIO": "Praia do Meio",
    "QUINTAS": "Quintas",
    "REDINHA": "Redinha",
    "RIBEIRA": "Ribeira",
    "ROCAS": "Rocas",
    "SALINAS": "Salinas",
    "SANTOS REIS": "Santos Reis",
    "TIROL": "Tirol",
}

NEIGHBORHOOD_ZONES = {
    "ALECRIM": "Leste",
    "AREIA PRETA": "Leste",
    "BARRO VERMELHO": "Leste",
    "BOM PASTOR": "Oeste",
    "CANDELARIA": "Sul",
    "CAPIM MACIO": "Sul",
    "CIDADE ALTA": "Leste",
    "CIDADE DA ESPERANCA": "Oeste",
    "CIDADE NOVA": "Oeste",
    "DIX SEPT ROSADO": "Oeste",
    "FELIPE CAMARAO": "Oeste",
    "GUARAPES": "Oeste",
    "IGAPO": "Norte",
    "LAGOA AZUL": "Norte",
    "LAGOA NOVA": "Sul",
    "LAGOA SECA": "Leste",
    "MAE LUIZA": "Leste",
    "MORRO BRANCO": "Sul",
    "NEOPOLIS": "Sul",
    "NORDESTE": "Oeste",
    "NOSSA SENHORA DA APRESENTACAO": "Norte",
    "NOSSA SENHORA DE NAZARE": "Oeste",
    "NOVA DESCOBERTA": "Sul",
    "PAJUCARA": "Norte",
    "PARQUE DAS COLINAS": "Sul",
    "PETROPOLIS": "Leste",
    "PITIMBU": "Sul",
    "PLANALTO": "Oeste",
    "PONTA NEGRA": "Sul",
    "POTENGI": "Norte",
    "PRAIA DO MEIO": "Leste",
    "QUINTAS": "Oeste",
    "REDINHA": "Norte",
    "RIBEIRA": "Leste",
    "ROCAS": "Leste",
    "SALINAS": "Norte",
    "SANTOS REIS": "Leste",
    "TIROL": "Leste",
}

NEIGHBORHOOD_RULES = (
    (r"APRESENT|\bN\s*S\s*A\b", "NOSSA SENHORA DA APRESENTACAO"),
    (r"NAZARE", "NOSSA SENHORA DE NAZARE"),
    (r"ALECR|AECRIM|ALEGRIM|ALENGRIM|ALERIM|ALCRIM|ALENCRIM", "ALECRIM"),
    (r"AREA PRETA", "AREIA PRETA"),
    (r"C\s*IDADE (?:DA )?E", "CIDADE DA ESPERANCA"),
    (r"^NOVA CIDADE$", "CIDADE NOVA"),
    (r"PET|PEPRO|P[T]?ROPOL", "PETROPOLIS"),
    (r"DIX.*SEPT|DIA SEPT|DIX DEPT|DIX SEP|DIX SERT|DIX SERPT|^SEPT ROSADO$", "DIX SEPT ROSADO"),
    (r"NEOP|NOEPO|NEOPO|NEO P", "NEOPOLIS"),
    (r"GUARAP", "GUARAPES"),
    (r"IGAP", "IGAPO"),
    (r"MAE LUIZA", "MAE LUIZA"),
    (r"FELIPE CAMAR|FILIPE CAMAR", "FELIPE CAMARAO"),
    (r"NOVA DESC|NOVA DECO|NOVA DESO|NOVA DESCOR", "NOVA DESCOBERTA"),
    (r"PAJUC|PJUC|PACUC|PACUJ|PAJAC|PAJC|PAJUAC|PAJUA|PAJUR", "PAJUCARA"),
    (r"LAGOA A?Z", "LAGOA AZUL"),
    (r"NOR.*DESTE|NODESTE|NOSDESTE|NOREDESTE|NOEDESTE", "NORDESTE"),
    (r"BARRO VERM|BAIRRO VERM", "BARRO VERMELHO"),
    (r"SANTO.*REI", "SANTOS REIS"),
    (r"QUINT|QUIT|QIU|QUNT", "QUINTAS"),
    (r"PRAI.*MEIO|PRIA.*MEIO|PRAIO.*MEIO", "PRAIA DO MEIO"),
    (r"POT|POTE|POTR|PONTENG", "POTENGI"),
    (r"PONTA.*NEG|PONTANEGRA|PONSTA NEG", "PONTA NEGRA"),
    (r"PIT|PINT", "PITIMBU"),
)


def text(value):
    if value is None or str(value).strip() == "":
        return "Nao informado"
    return str(value).strip()


def normalized(value):
    value = unicodedata.normalize("NFKD", text(value)).encode("ascii", "ignore").decode().lower()
    return " ".join("".join(char if char.isalnum() else " " for char in value).split())


def jaro_winkler(left, right):
    if left == right:
        return 1.0
    if not left or not right:
        return 0.0
    distance = max(len(left), len(right)) // 2 - 1
    left_matches = [False] * len(left)
    right_matches = [False] * len(right)
    matches = 0
    for index, char in enumerate(left):
        start = max(0, index - distance)
        end = min(index + distance + 1, len(right))
        for candidate in range(start, end):
            if right_matches[candidate] or char != right[candidate]:
                continue
            left_matches[index] = True
            right_matches[candidate] = True
            matches += 1
            break
    if not matches:
        return 0.0
    left_chars = [char for index, char in enumerate(left) if left_matches[index]]
    right_chars = [char for index, char in enumerate(right) if right_matches[index]]
    transpositions = sum(a != b for a, b in zip(left_chars, right_chars)) / 2
    jaro = (matches / len(left) + matches / len(right) + (matches - transpositions) / matches) / 3
    prefix = 0
    for a, b in zip(left[:4], right[:4]):
        if a != b:
            break
        prefix += 1
    return jaro + prefix * 0.1 * (1 - jaro)


def canonical_neighborhood(value):
    cleaned = normalized(value).upper()
    if not cleaned or cleaned == "NAO INFORMADO":
        return "Nao informado"
    if cleaned in OFFICIAL_NEIGHBORHOODS:
        return OFFICIAL_NEIGHBORHOODS[cleaned].upper()
    for pattern, official_key in NEIGHBORHOOD_RULES:
        if re.search(pattern, cleaned):
            return OFFICIAL_NEIGHBORHOODS[official_key].upper()
    best_key, similarity = max(
        ((candidate, jaro_winkler(cleaned, candidate)) for candidate in OFFICIAL_NEIGHBORHOODS),
        key=lambda item: item[1],
    )
    if similarity >= 0.85:
        return OFFICIAL_NEIGHBORHOODS[best_key].upper()
    return text(value).strip().upper()


def neighborhood_zone(value):
    return NEIGHBORHOOD_ZONES.get(normalized(value).upper(), "Nao informado")


def code(value):
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return text(value)


def load_decoders(data_dir, headers):
    workbook = openpyxl.load_workbook(
        os.path.join(data_dir, DICTIONARY_FILE), read_only=True, data_only=True
    )
    rows = workbook["dicionariotudo"].iter_rows(min_row=2, values_only=True)
    by_description = {}
    for _, description, answers in rows:
        if not description or not answers:
            continue
        mapping = {}
        for item in str(answers).split("#"):
            match = re.match(r"^\s*([^\-]+?)\s*-\s*(.+?)\s*$", item)
            if match:
                mapping[code(match.group(1))] = match.group(2).strip()
        if mapping:
            by_description[normalized(description)] = mapping
    return {
        index: by_description[normalized(header)]
        for index, header in enumerate(headers)
        if normalized(header) in by_description
    }


def decoded(decoders, index, value):
    if value is None:
        return "Nao informado"
    return decoders.get(index, {}).get(code(value), text(value))


def number(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def marked(value):
    value = normalized(value)
    return value == "1" or value.startswith("sim") or "opcao marcada" in value


def new_stats(filters):
    return {
        "filters": filters,
        "persons": 0,
        "families": 0,
        "familyPbf": 0,
        "incomeTotal": 0,
        "incomePerCapita": 0,
        "incomeCount": 0,
        "updated24": 0,
        "risk": 0,
        "foodRisk": 0,
        "indigenous": 0,
        "quilombola": 0,
        "personPbf": 0,
        "pcd": 0,
        "childLabor": 0,
        "children": 0,
        "street": 0,
        "gender": Counter(),
        "age": Counter(),
        "race": Counter(),
        "relationship": Counter(),
        "familyStatus": Counter(),
        "householdSize": Counter(),
        "updateMonths": Counter(),
        "housingType": Counter(),
        "rooms": Counter(),
        "sanitation": Counter(),
        "groups": Counter(),
        "disabilities": Counter(),
        "school": Counter(),
        "education": Counter(),
        "work": Counter(),
        "occupation": Counter(),
        "streetTime": Counter(),
        "streetSleep": Counter(),
        "services": Counter(),
        "expenses": Counter(),
        "expenseAmounts": {},
        "latestUpdate": None,
        "_families": set(),
    }


def add(counter, value):
    counter[text(value)] += 1


def serialize(stats):
    result = {}
    for key, value in stats.items():
        if key == "_families":
            continue
        result[key] = dict(value) if isinstance(value, Counter) else value
    return result


def post_with_retry(session, url, rows):
    payload = json.dumps(rows, ensure_ascii=False, allow_nan=False).encode("utf-8")
    last_error = None
    for attempt in range(1, 8):
        try:
            response = session.post(url, data=payload, timeout=120)
            if response.ok:
                return
            last_error = RuntimeError(f"HTTP {response.status_code}: {response.text}")
            if response.status_code < 500 and response.status_code != 429:
                raise last_error
        except (requests.ConnectionError, requests.Timeout) as error:
            last_error = error
        time.sleep(min(2**attempt, 30))
    raise last_error


def main():
    truststore.inject_into_ssl()
    base = os.environ["SUPABASE_URL"].rstrip("/")
    key = os.environ["SUPABASE_SECRET_KEY"]
    data_dir = os.path.join("src", "data")
    workbook = openpyxl.load_workbook(
        os.path.join(data_dir, SOURCE_FILE), read_only=True, data_only=True
    )
    sheet = workbook["tudo"]
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows)
    decoders = load_decoders(data_dir, headers)
    cube = {}

    for row_number, row in enumerate(rows, start=2):
        values = list(row) + [None] * (169 - len(row))
        localidade = canonical_neighborhood(decoded(decoders, 8, values[8]))
        filters = {
            "zona": neighborhood_zone(localidade),
            "localidade": localidade,
            "equipamento": decoded(decoders, 52, values[52]),
            "faixaRenda": decoded(decoders, 12, values[12]),
            "pbf": decoded(decoders, 14, values[14]),
        }
        key_tuple = tuple(filters.values())
        stats = cube.setdefault(key_tuple, new_stats(filters))
        stats["persons"] += 1
        stats["personPbf"] += int(marked(decoded(decoders, 73, values[73])))
        stats["pcd"] += int(marked(decoded(decoders, 79, values[79])))
        stats["childLabor"] += int(marked(decoded(decoders, 60, values[60])))
        stats["children"] += int(code(values[72]) in {"0", "1", "2", "3"})
        stats["street"] += int(marked(decoded(decoders, 121, values[121])))
        add(stats["gender"], decoded(decoders, 61, values[61]))
        add(stats["age"], decoded(decoders, 72, values[72]))
        add(stats["race"], decoded(decoders, 64, values[64]))
        add(stats["relationship"], decoded(decoders, 63, values[63]))
        add(stats["school"], decoded(decoders, 95, values[95]))
        add(stats["education"], decoded(decoders, 107, values[107]))
        add(stats["work"], decoded(decoders, 108, values[108]))
        add(stats["occupation"], decoded(decoders, 111, values[111]))

        if marked(decoded(decoders, 121, values[121])):
            add(stats["streetTime"], decoded(decoders, 130, values[130]))
            for index, label in ((122, "Rua"), (124, "Albergue"), (126, "Domicilio particular"), (128, "Outra forma")):
                if marked(decoded(decoders, index, values[index])):
                    stats["streetSleep"][label] += 1
            for index, label in ((150, "CRAS"), (151, "CREAS"), (152, "Centro POP"), (153, "Instituicao governamental"), (154, "Instituicao nao governamental"), (155, "Hospital ou clinica")):
                if marked(decoded(decoders, index, values[index])):
                    stats["services"][label] += 1

        disability_fields = (
            (80, "Cegueira"), (81, "Baixa visao"), (82, "Surdez severa ou profunda"),
            (83, "Surdez leve ou moderada"), (84, "Deficiencia fisica"),
            (85, "Deficiencia mental ou intelectual"), (86, "Sindrome de Down"),
            (87, "Transtorno ou doenca mental"),
        )
        for index, label in disability_fields:
            if marked(values[index]):
                stats["disabilities"][label] += 1

        family_code = text(values[1])
        if family_code not in stats["_families"]:
            stats["_families"].add(family_code)
            stats["families"] += 1
            stats["familyPbf"] += int(marked(decoded(decoders, 14, values[14])))
            stats["incomePerCapita"] += number(values[11])
            stats["incomeTotal"] += number(values[13])
            stats["incomeCount"] += 1
            months = number(values[15])
            stats["updated24"] += int(code(values[15]) in {"0", "1", "2"})
            stats["risk"] += int(marked(decoded(decoders, 54, values[54])))
            stats["foodRisk"] += int(marked(decoded(decoders, 55, values[55])))
            stats["indigenous"] += int(marked(decoded(decoders, 29, values[29])))
            stats["quilombola"] += int(marked(decoded(decoders, 35, values[35])))
            add(stats["familyStatus"], decoded(decoders, 4, values[4]))
            add(stats["householdSize"], values[38])
            add(stats["updateMonths"], decoded(decoders, 15, values[15]))
            add(stats["housingType"], decoded(decoders, 17, values[17]))
            add(stats["rooms"], values[18])
            add(stats["groups"], decoded(decoders, 56, values[56]))
            water = decoded(decoders, 22, values[22])
            bathroom = decoded(decoders, 24, values[24])
            sewage = decoded(decoders, 25, values[25])
            garbage = decoded(decoders, 26, values[26])
            lighting = decoded(decoders, 27, values[27])
            if marked(water): stats["sanitation"]["Agua canalizada"] += 1
            if marked(bathroom): stats["sanitation"]["Banheiro"] += 1
            if normalized(sewage) in {"rede coletora de esgoto ou pluvial", "fossa septica"}: stats["sanitation"]["Esgotamento adequado"] += 1
            if normalized(garbage) in {"e coletado diretamente", "e coletado indiretamente"}: stats["sanitation"]["Coleta de lixo"] += 1
            if normalized(lighting).startswith("eletrica"): stats["sanitation"]["Iluminacao eletrica"] += 1
            for index, label in ((43, "Energia"), (44, "Agua"), (45, "Gas"), (46, "Alimentacao"), (47, "Transporte"), (48, "Aluguel"), (49, "Medicamentos")):
                amount = number(values[index])
                if amount > 0:
                    stats["expenses"][label] += 1
                    amounts = stats["expenseAmounts"].setdefault(label, Counter())
                    amounts[f"{amount:g}"] += 1
            updated = values[3]
            if isinstance(updated, (date, datetime)):
                updated = updated.isoformat()
            elif updated:
                updated = str(updated)
            if updated and (not stats["latestUpdate"] or updated > stats["latestUpdate"]):
                stats["latestUpdate"] = updated

        if row_number % 50000 == 0:
            print(f"Linhas agregadas: {row_number:,}", flush=True)

    session = requests.Session()
    session.headers.update({
        "apikey": key,
        "Content-Type": "application/json",
        "Prefer": "resolution=merge-duplicates,return=minimal",
        "User-Agent": "cadunico-cube/1.0",
    })
    delete_url = f"{base}/rest/v1/excel_rows?source_file=eq.{CUBE_SOURCE}"
    response = session.delete(delete_url, timeout=120)
    response.raise_for_status()
    insert_url = f"{base}/rest/v1/excel_rows?on_conflict=source_file,sheet_name,row_number"
    records = [
        {
            "source_file": CUBE_SOURCE,
            "sheet_name": "filters",
            "row_number": index,
            "data": serialize(stats),
        }
        for index, stats in enumerate(cube.values(), start=1)
    ]
    for offset in range(0, len(records), BATCH_SIZE):
        post_with_retry(session, insert_url, records[offset : offset + BATCH_SIZE])
    print(f"Cubo enviado: {len(records):,} combinacoes", flush=True)


if __name__ == "__main__":
    main()
