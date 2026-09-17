import os
from openpyxl import load_workbook

base = r'c:\git\painel_cadunico\natal-social-lens\src\data'
for name in sorted(os.listdir(base)):
    if not name.lower().endswith(('.xlsx', '.xls')):
        continue
    path = os.path.join(base, name)
    print(f'\nFILE: {name}')
    wb = load_workbook(path, data_only=True)
    print('SHEETS:', wb.sheetnames)
    for ws in wb.worksheets[:3]:
        print(f'--- SHEET: {ws.title} rows={ws.max_row} cols={ws.max_column}')
        for row in ws.iter_rows(min_row=1, max_row=min(6, ws.max_row), values_only=True):
            print(row)
